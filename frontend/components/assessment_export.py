
from io import BytesIO
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
)


def get_question_text(question):
    return str(question.get("question", "Untitled question"))


def get_options(question):
    options = question.get("options") or []
    return options if isinstance(options, list) else []


def create_word_file(assessment, include_answers=False):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    title = assessment.get("assessment_title") or "Technical Assessment"
    document.add_heading(title, level=0)

    document.add_paragraph(
        f"Total Questions: {assessment.get('total_questions', 0)}"
    )
    document.add_paragraph(
        f"Total Marks: {assessment.get('total_marks', 0)}"
    )
    document.add_paragraph(
        f"Duration: {assessment.get('duration_minutes') or 'Not specified'} minutes"
    )

    document.add_heading("Instructions", level=1)
    document.add_paragraph(
        "Read each question carefully and answer according to the instructions."
    )

    document.add_heading("Questions", level=1)

    questions = assessment.get("questions", [])

    for index, question in enumerate(questions, start=1):
        document.add_paragraph(
            f"Q{index}. {get_question_text(question)} "
            f"({question.get('marks', 0)} marks)",
            style="List Number",
        )

        for option_index, option in enumerate(
            get_options(question), start=1
        ):
            document.add_paragraph(
                f"{chr(64 + option_index)}. {option}"
            )

        if question.get("question_type") == "coding":
            document.add_paragraph("Answer:")
            document.add_paragraph("\n\n")

        if include_answers and question.get("correct_answer"):
            document.add_paragraph(
                f"Answer key: {question['correct_answer']}"
            )

        if include_answers and question.get("explanation"):
            document.add_paragraph(
                f"Explanation: {question['explanation']}"
            )

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def create_pdf_file(assessment, include_answers=False):
    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="AssessmentTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#1F4E79"),
            spaceAfter=18,
        )
    )

    story = []

    title = escape(
        str(assessment.get("assessment_title") or "Technical Assessment")
    )

    story.append(Paragraph(title, styles["AssessmentTitle"]))

    story.append(
        Paragraph(
            f"Total Questions: {assessment.get('total_questions', 0)}",
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(
            f"Total Marks: {assessment.get('total_marks', 0)}",
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(
            f"Duration: {assessment.get('duration_minutes') or 'Not specified'} minutes",
            styles["Normal"],
        )
    )

    story.append(Spacer(1, 18))
    story.append(Paragraph("Instructions", styles["Heading2"]))
    story.append(
        Paragraph(
            "Read each question carefully and answer according to the instructions.",
            styles["Normal"],
        )
    )

    story.append(Spacer(1, 18))
    story.append(Paragraph("Questions", styles["Heading1"]))

    for index, question in enumerate(
        assessment.get("questions", []), start=1
    ):
        question_text = escape(get_question_text(question))
        marks = question.get("marks", 0)

        story.append(
            Paragraph(
                f"<b>Q{index}. {question_text}</b> ({marks} marks)",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 6))

        for option_index, option in enumerate(
            get_options(question), start=1
        ):
            option_text = escape(str(option))
            label = chr(64 + option_index)

            story.append(
                Paragraph(
                    f"&nbsp;&nbsp;{label}. {option_text}",
                    styles["Normal"],
                )
            )

        if question.get("question_type") == "coding":
            story.append(Spacer(1, 6))
            story.append(Paragraph("Answer:", styles["Normal"]))
            story.append(Spacer(1, 35))

        if include_answers and question.get("correct_answer"):
            answer = escape(str(question["correct_answer"]))
            story.append(
                Paragraph(f"<b>Answer key:</b> {answer}", styles["Normal"])
            )

        if include_answers and question.get("explanation"):
            explanation = escape(str(question["explanation"]))
            story.append(
                Paragraph(
                    f"<b>Explanation:</b> {explanation}",
                    styles["Normal"],
                )
            )

        story.append(Spacer(1, 14))

    pdf.build(story)
    output.seek(0)
    return output.getvalue()


def render_assessment_export():
    st.header("8. Export Assessment")

    assessment = st.session_state.get("assembled_assessment")

    if not assessment:
        st.info("Please assemble an assessment before exporting.")
        return

    st.write(
        "Download your assembled assessment as a Word document or PDF."
    )

    st.caption(
        f"{assessment.get('total_questions', 0)} questions · "
        f"{assessment.get('total_marks', 0)} marks"
    )

    include_answers = st.checkbox(
        "Include answer key and explanations",
        value=False,
        help="Keep this unchecked when sharing the assessment with candidates.",
    )

    title = assessment.get("assessment_title") or "technical_assessment"
    filename = "".join(
        character if character.isalnum() or character in "-_ " else "_"
        for character in title
    ).strip().replace(" ", "_")

    col1, col2 = st.columns(2)

    try:
        word_bytes = create_word_file(assessment, include_answers)
        pdf_bytes = create_pdf_file(assessment, include_answers)

        with col1:
            st.download_button(
                label="Download Word (.docx)",
                data=word_bytes,
                file_name=f"{filename}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

        with col2:
            st.download_button(
                label="Download PDF",
                data=pdf_bytes,
                file_name=f"{filename}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )

        st.success("Your assessment is ready to download.")

    except Exception as exc:
        st.error(f"Could not create the export files: {exc}")
